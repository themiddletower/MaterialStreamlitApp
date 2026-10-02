"""Convert the laboratory workbook to a validated table of individual repeats."""
import argparse
import hashlib
import json
import re
from numbers import Real
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from materials import PROPERTIES, ROOT


def read_fiber_sheet(path):
    book = openpyxl.load_workbook(path, data_only=True)
    matches = [s for s in book if 'ПТФЭ+УВ' in s.title.replace(' ', '')]
    if len(matches) != 1:
        raise ValueError('Ожидается один лист ПТФЭ+УВ.')
    sheet = matches[0]
    blocks = {}
    for row in sheet.iter_rows(min_row=1, max_row=5):
        for cell in row:
            text = str(cell.value or '').strip()
            for k, prefix in enumerate(['Предел прочности', 'Относительное удлинение',
                                         'Модуль упругости', 'Коэффициент трения', 'Износостойкость']):
                if text.startswith(prefix):
                    if k in blocks:
                        raise ValueError(f'Повторный заголовок: {prefix}')
                    blocks[k] = (cell.row, cell.column)
    if set(blocks) != set(range(5)):
        raise ValueError('На листе отсутствуют обязательные свойства.')
    # Each property is a separate block; never assume friction/wear column order.
    series = {}
    for k, (header, col) in blocks.items():
        modes = [str(sheet.cell(header + 1, col + m).value or '').strip() for m in range(3)]
        if modes != ['Статическое', 'УЗ+100', 'Статическое с активированным УВ']:
            raise ValueError(f'Неожиданные режимы для {PROPERTIES[k]}: {modes}')
        concentration, repeat = None, 0
        for r in range(header + 2, sheet.max_row + 1):
            label = sheet.cell(r, col - 1).value
            if label is not None:
                match = re.match(r'\s*(\d+(?:[.,]\d+)?)\s*%\s*УВ', str(label))
                concentration = float(match[1].replace(',', '.')) if match else None
                repeat = 0
            if concentration is None:
                continue
            values = [sheet.cell(r, col + m).value for m in range(3)]
            if not all(isinstance(v, Real) and not isinstance(v, bool) and np.isfinite(v) for v in values):
                raise ValueError(f'Пропуск или нечисловое измерение: {sheet.title}, строка {r}')
            repeat += 1
            for method, value in enumerate(values):
                series.setdefault((concentration, method, repeat), {})[PROPERTIES[k]] = float(value)
    rows = [dict(concentration=c, method=m, repeat=r, **values)
            for (c, m, r), values in sorted(series.items())]
    df = pd.DataFrame(rows)
    if df.empty or df[PROPERTIES].isna().any().any() or (df[PROPERTIES] <= 0).any().any():
        raise ValueError('Неполные или неположительные измерения.')
    counts = df.groupby(['concentration', 'method']).size()
    if len(counts) != 15 or not counts.eq(9).all() or sorted(df.concentration.unique()) != [10, 15, 20, 25, 30]:
        raise ValueError(f'Неожиданный план эксперимента: {counts.to_dict()}')
    metadata = {
        'source_file': Path(path).name,
        'sha256': hashlib.sha256(Path(path).read_bytes()).hexdigest(),
        'sheet': sheet.title, 'rows': len(df), 'groups': len(counts),
        'concentrations': sorted(df.concentration.unique().tolist()),
        'methods': modes,
        'notes': ['В заголовке относительного удлинения в исходнике указаны МПа; '
                  'в приложении используются проценты (как для одноимённого свойства УНТ). '
                  'Численные значения не изменены.',
                  'Строки средних исключены; чистый ПТФЭ и лист УНТ не включены в обучение УВ.'],
    }
    book.close()
    return df, metadata


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=Path)
    args = parser.parse_args()
    df, metadata = read_fiber_sheet(args.workbook)
    dest = ROOT / 'data'
    dest.mkdir(exist_ok=True)
    df.to_csv(dest / 'fiber_experiments.csv', index=False, encoding='utf-8-sig')
    (dest / 'fiber_source.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
