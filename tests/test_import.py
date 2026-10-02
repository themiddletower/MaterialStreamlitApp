import openpyxl
import pandas as pd
import pytest

from import_experiments import read_fiber_sheet
from materials import ROOT, PROPERTIES


def make_workbook(path, missing=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'ПТФЭ+УВ (3 режима)'
    data = pd.read_csv(ROOT / 'data/fiber_experiments.csv')
    # Deliberately put wear before friction, as in the supplied workbook.
    for block, k in enumerate([0,1,2,4,3]):
        col = 3 + block * 6
        sheet.cell(2,col,PROPERTIES[k])
        for m, name in enumerate(['Статическое','УЗ+100','Статическое с активированным УВ']):
            sheet.cell(3,col+m,name)
        for g, c in enumerate([10,15,20,25,30]):
            start = 4 + 10*g
            sheet.cell(start,col-1,f'{c}% УВ/ {100-c}% ПТФЭ')
            sheet.merge_cells(start_row=start, end_row=start+8, start_column=col-1, end_column=col-1)
            for r in range(9):
                for m in range(3):
                    value = float(data[(data.concentration==c)&(data.method==m)&(data.repeat==r+1)].iloc[0][PROPERTIES[k]])
                    sheet.cell(start+r,col+m,value)
            sheet.cell(start+9,col-1,'Среднее значение')
            for m in range(3):
                sheet.cell(start+9,col+m,999999)
    if missing:
        sheet.cell(5,4).value = None
    book.save(path)


def test_merged_headers_order_and_summary_rows(tmp_path):
    path = tmp_path / 'experiments.xlsx'
    make_workbook(path)
    actual, meta = read_fiber_sheet(path)
    expected = pd.read_csv(ROOT / 'data/fiber_experiments.csv')
    pd.testing.assert_frame_equal(actual[expected.columns], expected, check_dtype=False)
    assert meta['rows'] == 135 and meta['groups'] == 15


def test_missing_measurement_is_rejected(tmp_path):
    path = tmp_path / 'missing.xlsx'
    make_workbook(path, missing=True)
    with pytest.raises(ValueError, match='Пропуск'):
        read_fiber_sheet(path)
